"""현장 지시서: 발주서 초안(공급처·최소 발주량·단위 반영), 해동 지시, 엑셀·인쇄용 문서, 입고예정 추가용 행.

- 발주량은 '원료 소요·재고'의 발주 검토량에서 시작한다. 공급처의 최소 발주량보다 작으면 올리고, 발주 단위로 올림한다.
- 납품 요청일 = 긴급은 가장 빠른 입고일, 그 외는 발주 기한 + 리드타임(안전재고 아래로 내려가기 전에 들어와야 하는 날).
- 실제 발주를 보내지 않는다. 문서를 만들고, 완료 여부는 조치 이력에 사람이 기록한다.
"""
from __future__ import annotations

import html
import io
import math
import re

import pandas as pd

from lib import company as C
from lib import cost as K

COLUMNS = ["포함", "구분", "원료", "상태", "공급처", "발주량_kg", "계산량_kg", "납품요청일", "발주기한", "단가_원kg", "금액_원", "비고",
           "육종", "부위", "원산지", "보관"]


def supplier_for(data: C.CompanyData, key: tuple) -> dict:
    """원료의 공급처: 공급처.csv(우선순위가 작은 것) → 입고 예정의 가장 최근 공급처 → 미지정."""
    if data.supplier is not None and len(data.supplier):
        s = data.supplier[(data.supplier[C.MATERIAL_KEY] == pd.Series(key, index=C.MATERIAL_KEY)).all(axis=1)]
        if len(s):
            r = s.sort_values("우선순위").iloc[0]
            return {"공급처": r["공급처"], "담당자": r["담당자"], "연락처": r["연락처"],
                    "최소": float(r["최소발주량_kg"]), "단위": float(r["발주단위_kg"]), "근거": "공급처.csv"}
    if data.inbound is not None and "공급처" in data.inbound:
        i = data.inbound[(data.inbound[C.MATERIAL_KEY] == pd.Series(key, index=C.MATERIAL_KEY)).all(axis=1)]
        i = i[i["공급처"].notna() & ~i["공급처"].astype(str).isin(["", "nan"])]
        if len(i):
            return {"공급처": i.sort_values("입고예정일").iloc[-1]["공급처"], "담당자": "", "연락처": "", "최소": 0.0,
                    "단위": float(C.ROUND_KG), "근거": "최근 입고 예정"}
    return {"공급처": "미지정", "담당자": "", "연락처": "", "최소": 0.0, "단위": float(C.ROUND_KG), "근거": "없음"}


def round_order(qty: float, minimum: float, unit: float) -> tuple[float, str]:
    if qty <= 0:
        return 0.0, ""
    note = []
    q = qty
    if q < minimum:
        q = minimum
        note.append(f"최소 발주량 {minimum:,.0f}kg으로 올림")
    if unit > 0:
        r = math.ceil(q / unit - 1e-9) * unit
        if abs(r - q) > 1e-6:
            note.append(f"{unit:,.0f}kg 단위로 올림")
        q = r
    return q, ", ".join(note)


def draft(data: C.CompanyData, proj: C.Projection, settings: C.Settings | None = None) -> pd.DataFrame:
    """발주서 초안. 긴급·발주필요는 '포함'이 켜진 채로, 그 외 발주 검토량이 있는 원료는 꺼진 채로 넣는다."""
    st_ = settings or C.Settings()
    price = K.unit_price(K.material_prices(data), "next") if K.has_prices(data) else {}
    rows = []
    for _, r in proj.summary.iterrows():
        qty = float(r["발주량_kg"]) if proj.actionable else 0.0
        if qty <= 0:
            continue
        key = (r["육종"], r["부위"], r["원산지"], r["보관"])
        sup = supplier_for(data, key)
        q, note = round_order(qty, sup["최소"], sup["단위"])
        due = r["가장빠른입고일"] if r["상태"] == "긴급" or pd.isna(r["발주기한"]) else r["발주기한"] + pd.Timedelta(days=int(r["리드타임일수"]))
        p = price.get(r["원료"], (math.nan, ""))[0]
        rows.append({"포함": r["상태"] in ("긴급", "발주필요"), "구분": "정규 발주", "원료": r["원료"], "상태": r["상태"], "공급처": sup["공급처"],
                     "발주량_kg": q, "계산량_kg": qty, "납품요청일": due, "발주기한": r["발주기한"], "단가_원kg": p, "금액_원": q * p,
                     "비고": note, "육종": key[0], "부위": key[1], "원산지": key[2], "보관": key[3]})
    return pd.DataFrame(rows, columns=COLUMNS)


def from_inbound(data: C.CompanyData, rows: pd.DataFrame, kind: str = "긴급 현물(대응안)") -> pd.DataFrame:
    """대응안에서 추가한 입고 행(예: 긴급 현물)을 발주서 행으로 바꾼다."""
    out = []
    today = C.today_kst()
    for _, r in rows.iterrows():
        key = tuple(r[c] for c in C.MATERIAL_KEY)
        sup = supplier_for(data, key)
        q, note = round_order(float(r["입고량_kg"]), sup["최소"], sup["단위"])
        p = float(r[K.PRICE]) if K.PRICE in r and pd.notna(r[K.PRICE]) else math.nan
        out.append({"포함": True, "구분": kind, "원료": C.ingredient_label(*key), "상태": "대응안", "공급처": sup["공급처"] if sup["근거"] != "없음" else "긴급 현물 공급처(미지정)",
                    "발주량_kg": q, "계산량_kg": float(r["입고량_kg"]), "납품요청일": r["입고예정일"], "발주기한": today,
                    "단가_원kg": p, "금액_원": q * p, "비고": ", ".join(x for x in (note, "대응안 할증 단가(가정)") if x),
                    "육종": key[0], "부위": key[1], "원산지": key[2], "보관": key[3]})
    return pd.DataFrame(out, columns=COLUMNS)


def selected(order: pd.DataFrame) -> pd.DataFrame:
    o = order[order["포함"].fillna(False).astype(bool) & (order["발주량_kg"].fillna(0) > 0)].copy()
    o["금액_원"] = o["발주량_kg"] * o["단가_원kg"]
    return o


def inbound_rows(order: pd.DataFrame, today: pd.Timestamp | None = None) -> pd.DataFrame:
    """다음 업로드 때 입고예정.csv에 붙일 행 (원본은 자동으로 바꾸지 않는다)."""
    o = selected(order)
    d = pd.Timestamp(today or C.today_kst()).normalize()
    return pd.DataFrame({"입고예정일": pd.to_datetime(o["납품요청일"]).dt.strftime("%Y-%m-%d"), "육종": o["육종"], "부위": o["부위"],
                         "원산지": o["원산지"], "보관": o["보관"], "입고량_kg": o["발주량_kg"], "발주일": f"{d:%Y-%m-%d}",
                         "공급처": o["공급처"], "단가_원kg": o["단가_원kg"], "사용상태": "사용가능"})


def _sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", str(name))[:28] or "시트"
    out, i = base, 2
    while out in used:
        out, i = f"{base[:26]}_{i}", i + 1
    used.add(out)
    return out


def workbook(order: pd.DataFrame, thaw: pd.DataFrame, meta: dict) -> bytes:
    """공급처별 발주서 시트 + 해동 지시 + 입고예정 추가용 + 안내. openpyxl."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="C3C2B7")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="F3F2EE")
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()

    def table(ws, start_row: int, columns: list[str], rows: list[list], widths: list[int]):
        for j, c in enumerate(columns, 1):
            cell = ws.cell(start_row, j, c)
            cell.font, cell.fill, cell.border = Font(bold=True), head, box
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
        for i, r in enumerate(rows, start_row + 1):
            for j, v in enumerate(r, 1):
                cell = ws.cell(i, j, v)
                cell.border = box
                if isinstance(v, (int, float)):
                    cell.number_format = "#,##0"
        return start_row + len(rows) + 1

    sample = meta.get("샘플", False)
    o = selected(order)
    for supplier, g in o.groupby("공급처", sort=True):
        ws = wb.create_sheet(_sheet_name(f"발주_{supplier}", used))
        ws["A1"] = "발주서" + (" (샘플·가짜 데이터)" if sample else "")
        ws["A1"].font = Font(bold=True, size=16)
        info = [("발주처", supplier), ("발주일", meta.get("발주일", "")), ("발주 회사", meta.get("회사", "")), ("작성", meta.get("작성자", ""))]
        if len(g) and "담당자" in meta.get("공급처정보", {}).get(supplier, {}):
            si = meta["공급처정보"][supplier]
            info.append(("공급처 담당·연락처", f"{si.get('담당자', '')} {si.get('연락처', '')}".strip()))
        for i, (k, v) in enumerate(info, 3):
            ws.cell(i, 1, k).font = Font(bold=True)
            ws.cell(i, 2, v)
        rows = [[r["원료"], float(r["발주량_kg"]), f"{pd.Timestamp(r['납품요청일']):%Y-%m-%d}",
                 (float(r["단가_원kg"]) if pd.notna(r["단가_원kg"]) else None), (float(r["금액_원"]) if pd.notna(r["금액_원"]) else None),
                 r["구분"], r["비고"]] for _, r in g.iterrows()]
        end = table(ws, 3 + len(info) + 1, ["품목(육종 부위·원산지·보관)", "수량(kg)", "납품 요청일", "단가(원/kg)", "금액(원)", "구분", "비고"], rows,
                    [30, 12, 14, 13, 15, 16, 30])
        ws.cell(end, 1, "합계").font = Font(bold=True)
        ws.cell(end, 2, float(g["발주량_kg"].sum())).number_format = "#,##0"
        if g["금액_원"].notna().all():
            ws.cell(end, 5, float(g["금액_원"].sum())).number_format = "#,##0"
        ws.cell(end + 2, 1, "단가·금액은 참고값입니다. 최종 단가·납기는 공급처 확인 후 확정합니다.")
    ws = wb.create_sheet(_sheet_name("해동 지시", used))
    ws["A1"] = "해동 지시서" + (" (샘플·가짜 데이터)" if sample else "")
    ws["A1"].font = Font(bold=True, size=16)
    rows = [[f"{pd.Timestamp(r['해동시작일']):%Y-%m-%d}", r["원료"], r["로트"], round(float(r["냉동투입_kg"])), round(float(r["해동후_kg"])),
             f"{pd.Timestamp(r['사용일']):%Y-%m-%d}", r["상태"] if not r["대기일수"] else f"{r['상태']} · 해동 후 {int(r['대기일수'])}일 대기", "", ""]
            for _, r in thaw.iterrows()] if thaw is not None and len(thaw) else []
    table(ws, 3, ["해동 시작", "원료", "로트", "넣을 냉동 무게(kg)", "해동 후(kg)", "쓰는 날", "비고", "투입 확인", "확인자"], rows, [12, 26, 10, 16, 12, 12, 30, 10, 10])
    ws = wb.create_sheet(_sheet_name("입고예정 추가용", used))
    add = inbound_rows(order, pd.Timestamp(meta.get("발주일")) if meta.get("발주일") else None)
    table(ws, 1, list(add.columns), add.astype(object).where(add.notna(), None).values.tolist(), [12, 6, 8, 6, 6, 10, 12, 18, 10, 10])
    ws = wb.create_sheet(_sheet_name("안내", used))
    notes = ["이 파일은 대시보드가 계산한 검토용 문서입니다. 실제 발주·투입은 담당자가 확인한 뒤 진행합니다.",
             "발주량 = 발주 검토량을 공급처 최소 발주량·발주 단위에 맞춰 올린 값입니다.",
             "납품 요청일 = 긴급은 가장 빠른 입고일, 그 외는 발주 기한 + 리드타임입니다.",
             "'입고예정 추가용' 시트는 발주가 확정되면 입고예정.csv에 붙여 다음에 올리는 용도입니다(원본은 자동으로 바뀌지 않음).",
             "해동 지시의 '투입 확인'·'확인자' 칸은 현장에서 손으로 적거나, 대시보드 '현장 지시서'에서 체크해 조치 이력에 남깁니다."]
    if sample:
        notes.insert(0, "샘플(가짜) 데이터로 만든 문서입니다. 실제 발주에 쓰지 마세요.")
    for i, n in enumerate(notes, 1):
        ws.cell(i, 1, n)
    ws.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def print_html(order: pd.DataFrame, thaw: pd.DataFrame, meta: dict) -> str:
    """A4 인쇄용 지시서(HTML). 화면 안의 '인쇄' 버튼은 이 문서만 인쇄한다."""
    e = html.escape
    o = selected(order)
    sample = meta.get("샘플", False)

    def fmt(v, unit=""):
        return "-" if v is None or (isinstance(v, float) and math.isnan(v)) else f"{v:,.0f}{unit}"

    parts = [f"<h1>현장 지시서{' <small>(샘플·가짜 데이터)</small>' if sample else ''}</h1>",
             f"<p class='meta'>작성일 {e(str(meta.get('발주일', '')))} · 재고 기준일 {e(str(meta.get('기준일', '')))} · 작성 {e(str(meta.get('작성자', '')) or '-')}</p>"]
    parts.append("<h2>1. 발주</h2>")
    if len(o):
        for supplier, g in o.groupby("공급처", sort=True):
            parts.append(f"<h3>{e(str(supplier))}</h3><table><tr><th>품목</th><th>수량</th><th>납품 요청일</th><th>단가</th><th>금액</th><th>비고</th><th>확인</th></tr>")
            for _, r in g.iterrows():
                parts.append(f"<tr><td>{e(r['원료'])}</td><td class='n'>{fmt(r['발주량_kg'], 'kg')}</td><td>{pd.Timestamp(r['납품요청일']):%m/%d}</td>"
                             f"<td class='n'>{fmt(r['단가_원kg'], '원')}</td><td class='n'>{fmt(r['금액_원'], '원')}</td><td>{e(str(r['비고']))}</td><td class='chk'>□</td></tr>")
            parts.append(f"<tr class='sum'><td>합계</td><td class='n'>{fmt(float(g['발주량_kg'].sum()), 'kg')}</td><td></td><td></td>"
                         f"<td class='n'>{fmt(float(g['금액_원'].sum()) if g['금액_원'].notna().all() else math.nan, '원')}</td><td></td><td></td></tr></table>")
    else:
        parts.append("<p>포함된 발주가 없습니다.</p>")
    parts.append("<h2>2. 해동 투입</h2>")
    if thaw is not None and len(thaw):
        parts.append("<table><tr><th>해동 시작</th><th>원료</th><th>로트</th><th>넣을 냉동 무게</th><th>해동 후</th><th>쓰는 날</th><th>비고</th><th>투입 확인</th></tr>")
        for _, r in thaw.iterrows():
            memo = r["상태"] if not r["대기일수"] else f"{r['상태']} · 해동 후 {int(r['대기일수'])}일 대기"
            parts.append(f"<tr><td>{pd.Timestamp(r['해동시작일']):%m/%d}</td><td>{e(r['원료'])}</td><td>{e(str(r['로트']))}</td>"
                         f"<td class='n'>{fmt(float(r['냉동투입_kg']), 'kg')}</td><td class='n'>{fmt(float(r['해동후_kg']), 'kg')}</td>"
                         f"<td>{pd.Timestamp(r['사용일']):%m/%d}</td><td>{e(str(memo))}</td><td class='chk'>□</td></tr>")
        parts.append("</table>")
    else:
        parts.append("<p>잡힌 해동 투입이 없습니다.</p>")
    parts.append("<p class='foot'>대시보드 계산 결과로 만든 검토용 문서입니다. 단가·납기는 공급처 확인 후 확정하고, 실행 결과는 조치 이력에 남깁니다.</p>")
    css = """
    body{font-family:'Pretendard','Noto Sans KR','Malgun Gothic',sans-serif;color:#1a1a19;margin:16px}
    h1{font-size:20px;margin:0 0 4px} h1 small{font-size:13px;color:#d03b3b} h2{font-size:15px;margin:18px 0 6px} h3{font-size:13px;margin:10px 0 4px}
    .meta,.foot{font-size:12px;color:#52514e} table{border-collapse:collapse;width:100%;font-size:12px;margin-bottom:6px}
    th,td{border:1px solid #c3c2b7;padding:4px 6px;text-align:left} th{background:#f3f2ee} td.n{text-align:right;font-variant-numeric:tabular-nums}
    td.chk{text-align:center;width:48px} tr.sum td{font-weight:700} .bar{margin-bottom:10px}
    button{font:inherit;padding:6px 14px;border:1px solid #c3c2b7;background:#fff;border-radius:6px;cursor:pointer}
    @media print{.bar{display:none} body{margin:0} @page{size:A4;margin:12mm}}
    """
    return (f"<!doctype html><html lang='ko'><head><meta charset='utf-8'><title>현장 지시서</title><style>{css}</style></head><body>"
            "<div class='bar'><button onclick='window.print()'>인쇄</button></div>" + "".join(parts) + "</body></html>")


PENDING_TAG = "발주완료"  # 임시 반영한 입고 행의 로트 앞머리


def pending_inbound(events: list[dict], data: C.CompanyData, source: str) -> pd.DataFrame:
    """조치 이력의 '발주 완료' 기록 중 아직 입고예정.csv에 없는 물량. 원본 CSV는 바꾸지 않는다.

    - 같은 출처(샘플/올린 파일)의 기록만 쓴다. 사건마다 가장 최근 버전만 본다.
    - 업로드된 입고 예정에 같은 원료·입고예정일·수량(±0.5kg) 행이 있으면 이미 반영된 것으로 보고 뺀다.
    - 입고예정일이 재고 기준일 이전이면(이미 들어와 재고에 있어야 함) 뺀다.
    """
    latest: dict = {}
    for e in events:  # 최신 순
        latest.setdefault(e["사건ID"], e)
    base = data.inventory["기준일"].max()
    have = data.inbound if data.inbound is not None else pd.DataFrame(columns=C.REQUIRED["inbound"] + ["보관"])
    rows = []
    for case, e in latest.items():
        ctx = e.get("근거") or {}
        if ctx.get("종류") != "발주 완료" or ctx.get("출처") != source or e.get("상태") != "실행 완료":
            continue
        for i, r in enumerate(ctx.get("입고추가", [])):
            day = pd.Timestamp(r["입고예정일"])
            if day <= base:
                continue
            key = (r["육종"], r["부위"], r["원산지"], r["보관"])
            same = have[(have[C.MATERIAL_KEY] == pd.Series(key, index=C.MATERIAL_KEY)).all(axis=1)
                        & (have["입고예정일"] == day) & ((have["입고량_kg"] - float(r["입고량_kg"])).abs() < 0.5)]
            if len(same):
                continue
            rows.append({"입고예정일": day, "육종": key[0], "부위": key[1], "원산지": key[2], "보관": key[3],
                         "입고량_kg": float(r["입고량_kg"]), "발주일": pd.Timestamp(r.get("발주일")) if r.get("발주일") else pd.NaT,
                         "공급처": r.get("공급처", ""), "로트": f"{PENDING_TAG}-{case[:6]}-{i + 1}", "사용상태": "사용가능",
                         "사용가능일": day if key[3] == "냉장" else pd.NaT, "사용기한": pd.NaT,
                         **({K.PRICE: float(r[K.PRICE])} if r.get(K.PRICE) not in (None, "", "nan") and pd.notna(r.get(K.PRICE)) else {})})
    return pd.DataFrame(rows)


def apply_pending(data: C.CompanyData, rows: pd.DataFrame) -> C.CompanyData:
    """미반영 발주를 입고 예정에 더한 복사본."""
    from lib import planning as P
    if rows is None or rows.empty:
        return data
    inbound = pd.concat([data.inbound, rows], ignore_index=True) if data.inbound is not None else rows
    out = P.change_scenario(data, inbound=inbound)
    out.substitute, out.supplier = data.substitute, data.supplier
    return out


def is_pending_lot(lot_id) -> bool:
    return str(lot_id).startswith(PENDING_TAG + "-")
