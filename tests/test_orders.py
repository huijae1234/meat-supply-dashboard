"""현장 지시서: 공급처·최소 발주량·단위, 발주서 초안, 엑셀·인쇄 문서, 입고예정 추가용 행, 원산지 확인 항목, 화면."""
import io

import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from lib import advice as A
from lib import company as C
from lib import orders as O
from lib import response as R
from lib import thaw as T
from lib import ui_company as UC
from lib import alerts, live
from lib import worklog as W


def _sample():
    d = C.load_data(C.sample_tables())
    assert d.ready()
    return d, C.project(d, today=C.today_kst())


def test_round_order_minimum_and_unit():
    assert O.round_order(130, 500, 100) == (500, "최소 발주량 500kg으로 올림")
    q, note = O.round_order(1010, 0, 20)
    assert q == 1020 and "20kg 단위" in note
    assert O.round_order(0, 500, 100) == (0.0, "")


def test_supplier_fallbacks():
    d, _ = _sample()
    assert O.supplier_for(d, ("돼지", "뒷다리", "수입", "냉동"))["공급처"] == "B무역"
    d.supplier = None
    assert O.supplier_for(d, ("돼지", "앞다리", "국산", "냉장"))["근거"] == "최근 입고 예정"   # 입고예정의 A축산
    assert O.supplier_for(d, ("돼지", "갈비", "국산", "냉장"))["공급처"] == "미지정"


def test_supplier_table_validation():
    raw = C.sample_tables()
    raw["supplier"].loc[0, "공급처"] = ""
    assert not C.load_data(raw).ready()
    raw = C.sample_tables()
    raw["supplier"]["발주단위_kg"] = 0
    assert not C.load_data(raw).ready()
    raw = C.sample_tables()
    raw["supplier"] = raw["supplier"].drop(columns=["최소발주량_kg", "발주단위_kg", "담당자", "연락처", "우선순위"])
    d = C.load_data(raw)
    assert d.ready() and (d.supplier["발주단위_kg"] == C.ROUND_KG).all()


def test_draft_selects_urgent_and_due_date():
    d, p = _sample()
    o = O.draft(d, p)
    picked = o[o["포함"]]
    assert set(picked["상태"]) <= {"긴급", "발주필요"} and len(picked)
    assert (o.loc[~o["포함"], "상태"].isin(["긴급", "발주필요"]) == False).all()  # noqa: E712
    s = p.summary.set_index("원료")
    for _, r in o.iterrows():
        exp = s.loc[r["원료"], "가장빠른입고일"] if r["상태"] == "긴급" else s.loc[r["원료"], "발주기한"] + pd.Timedelta(days=int(s.loc[r["원료"], "리드타임일수"]))
        assert r["납품요청일"] == exp
        assert r["발주량_kg"] >= r["계산량_kg"]


def test_from_inbound_spot_rows():
    d, _ = _sample()
    rows = pd.DataFrame([{"입고예정일": pd.Timestamp("2026-10-20"), "육종": "돼지", "부위": "뒷다리", "원산지": "수입", "보관": "냉동",
                          "입고량_kg": 1010.0, "단가_원kg": 4830.0}])
    o = O.from_inbound(d, rows)
    assert o.iloc[0]["발주량_kg"] == 1020 and o.iloc[0]["공급처"] == "B무역" and o.iloc[0]["구분"] == "긴급 현물(대응안)"


def test_inbound_rows_can_be_uploaded_again():
    d, p = _sample()
    add = O.inbound_rows(O.draft(d, p))
    raw = C.sample_tables()
    raw["inbound"] = pd.concat([raw["inbound"], add], ignore_index=True)
    assert C.load_data(raw).ready()


def test_workbook_sheets_and_sample_mark():
    d, p = _sample()
    o = O.draft(d, p)
    o.loc[0, "공급처"] = "이름/에*특수"
    b = O.workbook(o, T.plan(d).events, {"발주일": "2026-10-18", "샘플": True})
    wb = openpyxl.load_workbook(io.BytesIO(b))
    assert {"해동 지시", "입고예정 추가용", "안내"} <= set(wb.sheetnames)
    assert any(n.startswith("발주_이름_에_특수") for n in wb.sheetnames)
    assert "샘플" in wb["안내"]["A1"].value
    ws = wb[[n for n in wb.sheetnames if n.startswith("발주_B")][0]]
    assert "샘플" in ws["A1"].value


def test_print_html_escapes_and_marks_sample():
    d, p = _sample()
    o = O.draft(d, p)
    o.loc[0, "비고"] = "<script>x</script>"
    page = O.print_html(o, T.plan(d).events, {"샘플": True, "작성자": "생산관리"})
    assert "<script>x</script>" not in page and "&lt;script&gt;" in page
    assert "샘플·가짜 데이터" in page and "window.print()" in page


def test_origin_changes_only_for_substitute():
    d = C.load_data(C.sample_tables())
    _, cands = R.candidates(d, today=C.today_kst())
    by = {c.name: R.origin_changes(d, c.data) for c in cands}
    assert by["대체 원료"] and not by["감산 수용"] and not by["긴급 현물 구매"]


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "COMPANY_DIR", tmp_path / "company")
    monkeypatch.delenv("COMPANY_PASSWORD", raising=False)
    monkeypatch.setattr(UC, "market_cached", lambda *args: A.Market())


def test_orders_page_sample_and_record(isolated):
    app = AppTest.from_string("from lib.ui_orders import render_orders\nrender_orders(1.5, lambda f, *a, **k: f)").run(timeout=60)
    app.button(key="orders_sample").click().run(timeout=60)
    assert not app.exception
    app.button(key="orders_done").click().run(timeout=60)
    assert any("기록하는 사람" in e.value for e in app.error)            # 기록자 없으면 막음
    app.text_input(key="orders_owner").set_value("생산관리 담당").run(timeout=60)
    app.button(key="orders_done").click().run(timeout=60)
    assert not app.exception
    ev = app.session_state["sample_worklog"][0]
    assert ev["상태"] == "실행 완료" and ev["제목"].startswith("발주 완료")
    app.checkbox(key="orders_spot").check().run(timeout=60)
    assert not app.exception
    assert not C.COMPANY_DIR.exists()                                      # 샘플은 디스크에 쓰지 않음


def test_substitute_scenario_brings_label_checklist_to_worklog(isolated):
    app = AppTest.from_string("import streamlit as st\nfrom lib.ui_planning import render_planning\nrender_planning()").run(timeout=60)
    app.button(key="planning_sample").click().run(timeout=60)
    app.selectbox[0].set_value("대체 원료").run(timeout=60)
    assert any("원산지 표시 확인" in w.value for w in app.warning)
    next(b for b in app.button if b.label == "이 대응안을 조치 기록에 연결").click().run(timeout=60)
    assert app.session_state["pending_action"]["확인항목"] == R.LABEL_CHECKS
    log = AppTest.from_string("import streamlit as st\nst.session_state['co_source']='sample'\n"
                              f"st.session_state['pending_action']={dict(app.session_state['pending_action'])!r}\n"
                              "from lib.ui_planning import render_worklog\nrender_worklog()").run(timeout=60)
    assert not log.exception
    assert any("□ 거래처(고객) 사전 승인" in t.value for t in log.text_area)


def _events(d, p, source="샘플(가짜) 데이터"):
    sel = O.selected(O.draft(d, p))
    add = O.inbound_rows(sel)
    return [{"사건ID": "abc123def", "상태": "실행 완료", "근거": {"종류": "발주 완료", "출처": source,
             "입고추가": add.astype(object).where(add.notna(), None).to_dict(orient="records")}}], sel


def test_pending_orders_are_counted_until_uploaded():
    d, p = _sample()
    events, sel = _events(d, p)
    rows = O.pending_inbound(events, d, "샘플(가짜) 데이터")
    assert len(rows) == len(sel) and rows["로트"].map(O.is_pending_lot).all()
    after = O.apply_pending(d, rows)
    assert after.supplier is not None and after.substitute is not None
    left = O.draft(after, C.project(after, today=C.today_kst()))
    assert not left["포함"].any()                                        # 같은 원료를 다시 발주하라고 하지 않음
    assert O.pending_inbound(events, d, "올려 둔 파일").empty           # 출처가 다르면 쓰지 않음
    raw = C.sample_tables()
    raw["inbound"] = pd.concat([raw["inbound"], O.inbound_rows(sel)], ignore_index=True)
    uploaded = C.load_data(raw)
    assert O.pending_inbound(events, uploaded, "샘플(가짜) 데이터").empty   # 업로드에 반영되면 자동으로 빠짐


def test_alerts_skip_ordered_materials(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "COMPANY_DIR", tmp_path / "company")
    monkeypatch.setattr(live, "LIVE_DIR", tmp_path / "live")
    for kind, frame in C.sample_tables().items():
        C.save_raw(kind, frame.to_csv(index=False).encode("utf-8-sig"))
    before = [e for e in alerts.find_alerts(alerts.CONFIG) if e["rule"] == "company" and "발주" in e["message"]]
    d = C.load_data(C.load_saved())
    events, sel = _events(d, C.project(d, today=C.today_kst()), "올려 둔 파일")
    ctx = events[0]["근거"]
    W.append_event("발주 완료", "생산관리", "2026-10-18", "실행 완료", "발주", "", ctx)
    after = [e for e in alerts.find_alerts(alerts.CONFIG) if e["rule"] == "company" and "발주" in e["message"]]
    assert len(after) < len(before)


def test_orders_record_then_today_line(isolated):
    app = AppTest.from_string("from lib.ui_orders import render_orders\nrender_orders(1.5, lambda f, *a, **k: f)").run(timeout=60)
    app.button(key="orders_sample").click().run(timeout=60)
    app.text_input(key="orders_owner").set_value("생산관리").run(timeout=60)
    app.button(key="orders_done").click().run(timeout=60)
    assert not app.exception
    assert any("(미반영)으로 계산에" in i.value for i in app.info)
    assert any(b.label.startswith("포함한 발주 0건") for b in app.button)
    events = app.session_state["sample_worklog"]
    today = AppTest.from_string("import streamlit as st\nst.session_state['co_source']='sample'\n"
                                f"st.session_state['sample_worklog']={events!r}\n"
                                "from lib.ui_company import render_today\nrender_today(1.5)").run(timeout=60)
    assert not today.exception
    assert any(m.value.startswith("**지시서**") and "발주 미기록 **0건**" in m.value for m in today.markdown)
